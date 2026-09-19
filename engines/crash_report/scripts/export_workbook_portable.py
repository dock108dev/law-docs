"""Linux workbook writer when the desktop spreadsheet runtime is unavailable."""

from datetime import date
from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


def export_workbook(model, out):
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows, widths):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        ws.append(headers)
        for row in rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    if cell.value.startswith(("=", "+", "@")):
                        cell.value = "'" + cell.value
                    cell.data_type = "s"
        for row in ws:
            for c in row:
                c.font = Font(name="Arial", size=11)
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor="243C50")
            c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        ws.row_dimensions[1].height = 34
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for i, row in enumerate(rows, 2):
            ws.row_dimensions[i].height = max(
                32,
                max(
                    (len(str(v or "")) // max(8, widths[j] - 3) + 1) * 15 + 10
                    for j, v in enumerate(row)
                ),
            )
        if rows:
            table = Table(
                displayName=name.replace(" ", "") + "Table",
                ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}",
            )
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        return ws

    reports = {r["id"]: r for r in model["reports"]}
    groups = model.get("households")
    if groups is not None:
        sheet(
            "Household mailings",
            [
                "Household",
                "Included People",
                "Street Address",
                "Mailing City",
                "State",
                "ZIP Code",
                "Report Pages File",
                "Report Cases",
                "Notes",
            ],
            [
                [
                    g["id"],
                    g["names"],
                    g["street_address"],
                    g["mailing_city"],
                    g["state"],
                    g["zip_code"],
                    g["id"] + "-report-pages.pdf",
                    "; ".join(reports[r]["case"] for r in g["reports"]),
                    "; ".join(g["notes"]),
                ]
                for g in groups
            ],
            [21, 42, 36, 24, 10, 18, 43, 28, 65],
        )
    keys = [
        "first_name",
        "last_name",
        "street_address",
        "mailing_city",
        "state",
        "zip_code",
        "crash_date",
        "crash_municipality",
    ]
    ws = sheet(
        "Recipients",
        [
            "First Name",
            "Last Name",
            "Street Address",
            "Mailing City",
            "State",
            "ZIP Code",
            "Crash Date",
            "Crash Municipality",
        ],
        [
            [date.fromisoformat(p[k]) if k == "crash_date" else p[k] for k in keys]
            for p in model["recipients"]
        ],
        [18, 26, 34, 24, 10, 18, 17, 24],
    )
    for row in ws.iter_rows(min_row=2):
        row[6].number_format = "mm/dd/yyyy"
    evidence = []
    for i, p in enumerate(model["recipients"], 1):
        r = reports[p["report"]]
        g = next((g for g in groups or [] if any(x["id"] == p["id"] for x in g["people"])), None)
        evidence.append(
            [
                i,
                p["first_name"] + " " + p["last_name"],
                r["case"],
                r["department"],
                p["vehicle"],
                p["role"],
                p["selection"],
                ", ".join(map(str, r["source_pages"])),
                p["first_page"],
                g["reports"].index(p["report"]) + 1 if g else i,
                ", ".join(map(str, p["support_pages"])),
                g["id"] + "-report-pages.pdf" if g else model.get("correction_status", "Automatic"),
            ]
        )
    sheet(
        "Recipient evidence",
        [
            "Recipient Row",
            "Recipient",
            "Report Case",
            "Police Department",
            "Vehicle",
            "Role",
            "Selection Evidence",
            "Source Pages",
            "First Page",
            "Packet Page" if groups is not None else "Proof Page",
            "Support Pages",
            "Household File" if groups is not None else "Extraction",
        ],
        evidence,
        [15, 29, 23, 28, 12, 14, 55, 16, 14, 14, 16, 43],
    )
    sheet(
        "Reports",
        [
            "Report",
            "Case",
            "Police Department",
            "Crash Date",
            "Municipality",
            "Source Pages",
            "First Page",
            "118 A/B",
            "119 A/B",
            "Review Notes",
        ],
        [
            [
                r["id"],
                r["case"],
                r["department"],
                r["crash_date"],
                r["municipality"],
                ", ".join(map(str, r["source_pages"])),
                r["first_page"],
                r["codes"]["118a"] + "/" + r["codes"]["118b"],
                r["codes"]["119a"] + "/" + r["codes"]["119b"],
                "; ".join(r["issues"]),
            ]
            for r in model["reports"]
        ],
        [12, 23, 28, 16, 24, 16, 14, 16, 16, 95],
    )
    rows = []
    for r in model["review"]:
        p = r.get("candidate", {})
        f = r.get("fields", {})
        rows.append(
            [
                r.get("report", p.get("report", "")),
                r.get("occupant_row", ""),
                (p.get("first_name", "") + " " + p.get("last_name", "")).strip() or f.get("95", ""),
                r.get("role", p.get("role", "")),
                "; ".join(k + "=" + f.get(k, "") for k in ["83", "84", "86"]) if f else "",
                r.get("source_page", p.get("first_page", "")),
                r["reason"],
            ]
        )
    sheet(
        "Review",
        [
            "Report",
            "Occupant",
            "Person / Source Text",
            "Role",
            "Occupant Codes",
            "Source Page",
            "Disposition / Reason",
        ],
        rows,
        [12, 12, 90, 16, 34, 15, 80],
    )
    if "drivers" in model:
        keys = [
            "id",
            "source_page",
            "role",
            "vehicle",
            "first_name",
            "middle",
            "last_name",
            "street_address",
            "mailing_city",
            "state",
            "zip_code",
        ]
        sheet(
            "Extracted drivers",
            [
                "ID",
                "Source Page",
                "Role",
                "Vehicle",
                "First Name",
                "Middle",
                "Last Name",
                "Street",
                "City",
                "State",
                "ZIP",
                "Contributing Codes",
                "Disposition",
            ],
            [
                [r.get(k, "") for k in keys] + ["/".join(r["codes"]), r["disposition"]]
                for r in model["drivers"]
            ],
            [24, 14, 34, 12, 20, 10, 28, 36, 24, 10, 18, 24, 75],
        )
    if "occupants" in model:
        keys = [
            "id",
            "source_page",
            "occupant_row",
            "vehicle",
            "seat",
            "injury",
            "first_name",
            "last_name",
            "street_address",
            "mailing_city",
            "state",
            "zip_code",
            "raw_name_address",
            "parse_issue",
        ]
        sheet(
            "Extracted occupants",
            [
                "ID",
                "Source Page",
                "Row",
                "Vehicle",
                "Seat",
                "Injury",
                "First Name",
                "Last Name",
                "Street",
                "City",
                "State",
                "ZIP",
                "Raw Box 95",
                "Parse Note",
            ],
            [[r.get(k, "") for k in keys] for r in model["occupants"]],
            [24, 14, 10, 12, 10, 10, 20, 28, 36, 24, 10, 18, 110, 50],
        )
    wb.save(Path(out) / "recipients.xlsx")
