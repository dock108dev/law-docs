"""One mailing group per complete address; retain every selected person."""

import re
import pymupdf as fitz
import openpyxl


def address_key(person):
    def normalized(value):
        return " ".join(re.sub(r"[.,]", "", str(value).upper()).split())

    parts = [
        normalized(person.get(k, ""))
        for k in ("street_address", "mailing_city", "state", "zip_code")
    ]
    if not all(parts):
        return ("unresolved", person["id"])
    # Normalize common street/unit spelling, but never discard unit numbers.
    aliases = {
        "STREET": "ST",
        "AVENUE": "AVE",
        "ROAD": "RD",
        "DRIVE": "DR",
        "LANE": "LN",
        "BOULEVARD": "BLVD",
        "COURT": "CT",
        "PLACE": "PL",
        "APARTMENT": "UNIT",
        "APT": "UNIT",
        "SUITE": "UNIT",
        "STE": "UNIT",
    }
    street = re.sub(r"#\s*", " UNIT ", parts[0])
    parts[0] = " ".join(aliases.get(word, word) for word in street.split())
    parts[3] = parts[3][:5] if re.fullmatch(r"\d{5}(?:-\d{4})?", parts[3]) else parts[3]
    return tuple(parts)


def group_households(people):
    groups = {}
    for person in people:
        key = address_key(person)
        group = groups.setdefault(
            key,
            {
                "id": f"household-{len(groups) + 1:03}",
                "street_address": person.get("street_address", ""),
                "mailing_city": person.get("mailing_city", ""),
                "state": person.get("state", ""),
                "zip_code": person.get("zip_code", ""),
                "people": [],
                "reports": [],
                "notes": [],
            },
        )
        group["people"].append(person)
        if person["report"] not in group["reports"]:
            group["reports"].append(person["report"])
        if key[0] == "unresolved":
            group["notes"].append("Incomplete address; check before mailing. Kept separate.")
        if person.get("manual_exception"):
            group["notes"].append(
                "Manual inclusion: " + person.get("exception_reason", "Previously approved")
            )
    for group in groups.values():
        group["names"] = "; ".join(
            dict.fromkeys(
                " ".join(filter(None, [p.get("first_name"), p.get("middle"), p.get("last_name")]))
                for p in group["people"]
            )
        )
        group["notes"] = list(dict.fromkeys(group["notes"]))
    return list(groups.values())


def export_household_pages(model, out):
    """One original first page per household/report; no repeated family copies."""
    mapping = []
    with fitz.open(model["source_pdf"]) as source, fitz.open() as combined:
        for group in model["households"]:
            with fitz.open() as packet:
                for report_id in group["reports"]:
                    report = next(r for r in model["reports"] if r["id"] == report_id)
                    index = report["first_page"] - 1
                    packet.insert_pdf(source, from_page=index, to_page=index)
                    combined.insert_pdf(source, from_page=index, to_page=index)
                    mapping.append(
                        {
                            "household_id": group["id"],
                            "packet_file": group["id"] + "-report-pages.pdf",
                            "packet_page": len(packet),
                            "combined_page": len(combined),
                            "report": report_id,
                            "case": report["case"],
                            "source_page": index + 1,
                            "people": [
                                p["id"] for p in group["people"] if p["report"] == report_id
                            ],
                        }
                    )
                packet.save(out / (group["id"] + "-report-pages.pdf"))
        if len(combined):
            combined.save(out / "household-report-pages.pdf")
    # Reopen and compare every delivered first page with its original rendering.
    for entry in mapping:
        with (
            fitz.open(out / entry["packet_file"]) as packet,
            fitz.open(model["source_pdf"]) as source,
        ):
            assert (
                packet[entry["packet_page"] - 1].get_pixmap().samples
                == source[entry["source_page"] - 1].get_pixmap().samples
            )
    return mapping


def verify_household_workbook(model, out, mapping):
    workbook = openpyxl.load_workbook(out / "recipients.xlsx", data_only=False)
    sheet = workbook["Household mailings"]
    assert sheet.max_row == len(model["households"]) + 1
    assert workbook.sheetnames[0] == "Household mailings"
    for index, group in enumerate(model["households"], 2):
        values = [
            group["id"],
            group["names"],
            group["street_address"],
            group["mailing_city"],
            group["state"],
            group["zip_code"],
            group["id"] + "-report-pages.pdf",
        ]
        for column, value in enumerate(values, 1):
            # Exporter escapes formula-like source strings as text.
            expected = "'" + value if value.startswith(("=", "+", "@")) else value
            assert (sheet.cell(index, column).value or "") == expected
        assert len([p for p in mapping if p["household_id"] == group["id"]]) == len(
            group["reports"]
        )
    workbook.close()
    return {
        "households": len(model["households"]),
        "included_people": len(model["recipients"]),
        "distinct_household_report_pages": len(mapping),
        "mailing_rows_verified": True,
        "original_first_page_pixels_verified": True,
    }
