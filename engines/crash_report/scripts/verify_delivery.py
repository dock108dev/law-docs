#!/usr/bin/env python3
"""Verify each serialized workbook row and proof page against the extraction list."""

import argparse, hashlib, json
from pathlib import Path
from source_paths import source_path
from safe_paths import contained
from datetime import datetime
import openpyxl
import pymupdf as fitz
import numpy as np


def verify(out, raw):
    out = contained(out)
    m = json.loads(contained(out / "final.json", out).read_text())
    manifest = json.loads(contained(out / "proof-manifest.json", out).read_text())
    w = openpyxl.load_workbook(contained(out / "recipients.xlsx", out), data_only=False)
    sheet = w["Recipients"]
    keys = [
        "first_name",
        "last_name",
        "street_address",
        "mailing_city",
        "state",
        "zip_code",
        "crash_date",
        "crash_municipality",
    ]
    assert sheet.max_column == 8
    assert (
        hashlib.sha256(source_path(m["source_pdf"]).read_bytes()).hexdigest() == m["source_sha256"]
    )
    checks = []
    with (
        fitz.open(contained(out / "recipient-proofs.pdf", out)) as proof,
        fitz.open(source_path(m["source_pdf"])) as source,
    ):
        assert sheet.max_row - 1 == len(proof) == len(manifest) == len(m["recipients"])
        for i, (r, p, entry) in enumerate(zip(m["recipients"], proof, manifest), 1):
            for col, key in enumerate(keys, 1):
                cell = sheet.cell(i + 1, col)
                if key == "crash_date":
                    assert (
                        isinstance(cell.value, datetime) and cell.value.date().isoformat() == r[key]
                    )
                else:
                    assert cell.value == r[key]
            assert sheet.cell(i + 1, 6).data_type == "s"
            assert entry["recipient_id"] == r["id"]
            assert entry["label"] == f"Recipient row {i} | {r['first_name']} {r['last_name']}"
            assert entry["label"] in p.get_text()
            report = next(x for x in m["reports"] if x["id"] == r["report"])
            assert entry["source_page"] == report["first_page"] == r["first_page"]
            assert entry["proof_page"] == entry["recipient_row"] == i
            assert [x["field"] for x in entry["highlights"]] == r["evidence_fields"]
            drawings = p.get_drawings()
            for highlight in entry["highlights"]:
                expected = raw["pages"][r["first_page"] - 1]["fields"][highlight["field"]]["rect"]
                assert highlight["source_rect"] == expected
                shifted = fitz.Rect(expected) + fitz.Rect(0, 48, 0, 48)
                assert any(
                    max(abs(a - b) for a, b in zip(d["rect"], shifted)) < 0.01 for d in drawings
                )
            original = source[r["first_page"] - 1].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            actual = p.get_pixmap(
                matrix=fitz.Matrix(2, 2),
                clip=fitz.Rect(0, 48, p.rect.width, p.rect.height),
                alpha=False,
            )
            a = np.frombuffer(original.samples, dtype=np.uint8).reshape(
                original.height, original.width, 3
            )
            b = np.frombuffer(actual.samples, dtype=np.uint8).reshape(
                actual.height, actual.width, 3
            )
            assert a.shape == b.shape
            mask = np.ones(a.shape[:2], dtype=bool)
            for h in entry["highlights"]:
                x0, y0, x1, y1 = h["source_rect"]
                mask[
                    max(0, int(y0 * 2) - 3) : int(y1 * 2) + 4,
                    max(0, int(x0 * 2) - 3) : int(x1 * 2) + 4,
                ] = False
            difference = float(np.abs(a[mask].astype(float) - b[mask].astype(float)).mean())
            assert difference < 0.1
            checks.append(
                {
                    "row": i,
                    "recipient_id": r["id"],
                    "name": entry["name"],
                    "case": report["case"],
                    "source_page": r["first_page"],
                    "proof_page": i,
                    "source_image_mean_difference_outside_highlights": difference,
                    "result": "pass",
                }
            )
    return {
        "rows": checks,
        "excel_rows": len(checks),
        "proof_pages": len(checks),
        "source_hash_unchanged": True,
        "visual_review": "Recorded separately; programmatic checks do not imply visual or owner acceptance.",
    }


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("directory", type=Path)
    ap.add_argument("--raw", type=Path, required=True)
    ap.add_argument("--result", type=Path, required=True)
    a = ap.parse_args()
    result = verify(a.directory, json.loads(a.raw.read_text()))
    a.result.write_text(json.dumps(result, indent=2))
    print(
        f"Verified {result['excel_rows']} individual Excel/proof pairs, coordinates and first-page pixels."
    )
